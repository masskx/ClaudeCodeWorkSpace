#!/usr/bin/env python3
"""按 notes.json 里的读书笔记，生成一份 10 分钟组会 PPT（组会示例.pptx）。
用法：python3 make_ppt.py    需要 python-pptx。
结构：封面 → 先说结论 → 为什么是这五篇 → 五篇各一页 → 放在一起看 → 对我们课题的启发 → 备用页（老师可能会问）。
页面里的图都是用形状自己画的，没有复制论文里的图；你自己汇报时如果要放原图，请标明出处。"""
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
N = json.load(open(os.path.join(HERE, "notes.json"), encoding="utf-8"))
FONT = "Microsoft YaHei"
INK, MUTE, LINE = RGBColor(0x1f,0x23,0x2b), RGBColor(0x6b,0x72,0x80), RGBColor(0xe3,0xe6,0xec)
BLUE, ORANGE, SOFT, SOFT2 = RGBColor(0x2a,0x78,0xd6), RGBColor(0xeb,0x68,0x34), RGBColor(0xee,0xf4,0xfc), RGBColor(0xf5,0xf6,0xf8)
WHITE = RGBColor(255,255,255)

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]

def font(run, size, bold=False, color=INK):
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, FONT
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", FONT)

def text(slide, x, y, w, h, paras, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    if isinstance(paras, str): paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.space_after = Pt(space)
        r = para.add_run(); r.text = p; font(r, size, bold, color)
    return tb

def box(slide, x, y, w, h, fill=SOFT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    s.shadow.inherit = False
    st = s._element.find(qn('p:style'))
    if st is not None: s._element.remove(st)
    return s

def label(slide, x, y, w, h, t, size=13, fill=SOFT, color=INK, bold=False, line=None):
    s = box(slide, x, y, w, h, fill=fill, line=line)
    tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t; font(r, size, bold, color)
    return s

def header(slide, tag, title, sub=None):
    text(slide, 0.6, 0.35, 9, 0.35, tag, 13, False, MUTE)
    text(slide, 0.6, 0.65, 12.1, 0.8, title, 30, True, INK)
    if sub: text(slide, 0.6, 1.4, 12.1, 0.4, sub, 13, False, MUTE)

def footer(slide, t):
    box(slide, 0.6, 6.55, 12.13, 0.65, fill=BLUE)
    text(slide, 0.8, 6.58, 11.8, 0.6, t, 17, True, WHITE, anchor=MSO_ANCHOR.MIDDLE)

def bullets(slide, x, y, w, items, size=16, color=INK, gap=6):
    return text(slide, x, y, w, 0.5*len(items)+0.4, ["• " + i for i in items], size, False, color, space=gap)

# ---------- diagrams (自己画的示意，不是论文原图) ----------
def d_resnet(s, x, y):
    label(s, x+0.05, y, 0.55, 0.5, "x", 15, SOFT, bold=True)
    text(s, x+0.6, y+0.02, 0.3, 0.4, "→", 16)
    label(s, x+0.9, y-0.02, 1.2, 0.55, "权重层 F(x)", 12, SOFT)
    text(s, x+2.1, y+0.02, 0.3, 0.4, "→", 16)
    label(s, x+2.4, y, 0.45, 0.5, "+", 16, ORANGE, WHITE, True)
    text(s, x+2.85, y+0.02, 0.3, 0.4, "→", 16)
    label(s, x+3.15, y, 1.05, 0.5, "F(x) + x", 13, SOFT, bold=True)
    text(s, x+0.05, y+0.75, 4.2, 0.4, "↳ 快捷连接：把 x 直接加到输出", 13, False, ORANGE)
    text(s, x+0.05, y+1.2, 4.2, 0.4, "加了快捷连接，参数量不变", 12, False, MUTE)
def d_transformer(s, x, y):
    for i, t in enumerate(["前馈层", "多头自注意力", "输入 + 位置编码"]):
        label(s, x+0.2, y+i*0.7, 3.0, 0.55, t, 13, SOFT if i else ORANGE, WHITE if i == 0 else INK, bool(i == 0))
    text(s, x+3.3, y+0.65, 0.9, 0.5, "× N 层", 14, True, MUTE)
    text(s, x+0.2, y+2.0, 4.0, 0.4, "没有循环，没有卷积", 12, False, MUTE)
def d_vit(s, x, y):
    for r in range(3):
        for c in range(3):
            box(s, x+0.1+c*0.42, y+r*0.42, 0.38, 0.38, fill=SOFT if (r+c) % 2 else RGBColor(0xd9,0xe8,0xfa), shape=MSO_SHAPE.RECTANGLE)
    text(s, x+1.45, y+0.4, 0.3, 0.4, "→", 16)
    for i, t in enumerate(["CLS", "p1", "p2", "p3"]):
        label(s, x+1.8+i*0.6, y+0.4, 0.52, 0.45, t, 11, ORANGE if i == 0 else SOFT, WHITE if i == 0 else INK)
    label(s, x+1.8, y+1.05, 2.4, 0.5, "Transformer 编码器", 12, SOFT, bold=True)
    text(s, x+0.05, y+1.35, 1.6, 0.4, "切成 patch", 12, False, MUTE)
def d_clip(s, x, y):
    label(s, x+0.05, y, 1.3, 0.45, "图像编码器", 12, SOFT); label(s, x+0.05, y+0.6, 1.3, 0.45, "文本编码器", 12, SOFT)
    text(s, x+1.4, y+0.4, 0.3, 0.4, "→", 16)
    for r in range(4):
        for c in range(4):
            box(s, x+1.8+c*0.5, y+r*0.42, 0.45, 0.38, fill=ORANGE if r == c else SOFT2, shape=MSO_SHAPE.RECTANGLE)
    text(s, x+0.05, y+1.8, 4.2, 0.4, "对角线 = 配对的图文；其余都是负例", 12, False, MUTE)
def d_ddpm(s, x, y):
    greys = ["333333", "666666", "999999", "bbbbbb", "dddddd"]
    for i, g in enumerate(greys):
        box(s, x+0.1+i*0.82, y, 0.62, 0.62, fill=RGBColor.from_string(g), shape=MSO_SHAPE.RECTANGLE)
        text(s, x+0.1+i*0.82, y+0.64, 0.62, 0.3, ["x0", "x1", "…", "…", "xT"][i], 12, False, MUTE, PP_ALIGN.CENTER)
    text(s, x+0.1, y+1.0, 4.2, 0.4, "→ 逐步加噪（前向）", 13, False, MUTE)
    text(s, x+0.1, y+1.4, 4.2, 0.4, "← 学习逐步去噪（生成）", 13, True, ORANGE)
DIAG = {"resnet": d_resnet, "transformer": d_transformer, "vit": d_vit, "clip": d_clip, "ddpm": d_ddpm}

# ---------- 1 封面 ----------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, 13.333, 7.5, fill=SOFT, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.9, 2.2, 11.5, 1.2, N["meeting"]["title"], 48, True, INK)
text(s, 0.9, 3.5, 11.5, 0.6, N["meeting"]["sub"], 24, False, BLUE)
text(s, 0.9, 5.6, 11.5, 0.5, N["meeting"]["who"], 16, False, MUTE)
text(s, 0.9, 6.1, 11.5, 0.5, "10 分钟 · 5 篇论文", 14, False, MUTE)

# ---------- 2 先说结论 ----------
s = prs.slides.add_slide(BLANK)
header(s, "先说结论（30 秒）", "这五篇讲的是同一件事：谁来提供“视觉的先验”")
cards = [("1", "结构先验让位给规模", "ResNet 解决“深了训不动”；ViT 几乎不加图像先验，靠大规模预训练追平并超过 CNN。"),
         ("2", "监督信号从标签换成语言", "CLIP 用 4 亿图文对预训练，用文字当接口，零样本做分类。"),
         ("3", "生成是另一条线", "DDPM 用逐步去噪生成图像，样本质量高、训练稳，代价是采样步数多。")]
for i, (n, h, b) in enumerate(cards):
    x = 0.6 + i*4.1
    box(s, x, 1.9, 3.9, 3.9, fill=SOFT2)
    label(s, x+0.25, 2.15, 0.6, 0.6, n, 20, BLUE, WHITE, True)
    text(s, x+0.25, 2.95, 3.4, 0.9, h, 22, True, INK)
    text(s, x+0.25, 3.95, 3.4, 1.8, b, 18, False, MUTE)
footer(s, "对我们课题：先把“强预训练模型 + 线性探测”跑出来当基线，再谈改进。")

# ---------- 3 为什么是这五篇 ----------
s = prs.slides.add_slide(BLANK)
header(s, "为什么是这五篇", "按时间看，这五篇各占一个位置")
line = box(s, 0.9, 3.55, 11.5, 0.06, fill=LINE, shape=MSO_SHAPE.RECTANGLE)
DATES = {"resnet": "2015.12", "transformer": "2017.06", "ddpm": "2020.06", "vit": "2020.10", "clip": "2021.02"}
order = sorted(N["papers"], key=lambda p: DATES[p["key"]])
for i, p in enumerate(order):
    x = 0.9 + i*2.55
    box(s, x+0.7, 3.4, 0.35, 0.35, fill=ORANGE if p["key"] == "vit" else BLUE, shape=MSO_SHAPE.OVAL)
    text(s, x-0.1, 2.7, 1.95, 0.5, DATES[p["key"]], 20, True, INK, PP_ALIGN.CENTER)
    text(s, x-0.15, 3.95, 2.05, 0.5, p["head"].split("：")[0], 20, True, BLUE, PP_ALIGN.CENTER)
    text(s, x-0.2, 4.5, 2.15, 1.3, p["head"].split("：")[1], 13, False, MUTE, PP_ALIGN.CENTER)
text(s, 0.9, 5.9, 11.5, 0.4, "时间取自论文在 arXiv 的首版日期（年月）。汇报顺序是 ResNet → Transformer → ViT → CLIP → DDPM，便于讲清“视觉骨干”这条主线。", 12, False, MUTE)
footer(s, "怎么读：每篇一页，问题 → 做法 → 最关键的证据 → 作者承认的局限。")

# ---------- 4–8 五篇各一页 ----------
for p in N["papers"]:
    s = prs.slides.add_slide(BLANK)
    header(s, f"论文 {p['n']} / 5 · {p['year']}", p["head"], f"{p['cite']}（arXiv:{p['arxiv']}）")
    text(s, 0.6, 1.95, 4.0, 0.35, "要解决什么", 14, True, BLUE)
    bullets(s, 0.6, 2.3, 4.0, p["problem"], 17)
    text(s, 0.6, 3.95, 4.0, 0.35, "怎么做", 15, True, BLUE)
    bullets(s, 0.6, 4.35, 4.0, p["method"], 17)
    box(s, 4.85, 1.95, 4.4, 2.65, fill=SOFT2); DIAG[p["key"]](s, 5.0, 2.25)
    box(s, 9.5, 1.95, 3.23, 1.85, fill=SOFT)
    text(s, 9.62, 2.0, 3.0, 0.8, p["stat"], 30 if len(p["stat"]) <= 9 else 22, True, BLUE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.62, 2.85, 3.0, 0.9, p["stat_cap"], 13, False, MUTE)
    text(s, 4.85, 4.75, 4.4, 1.7, [p["extra"]], 15, False, INK)
    text(s, 9.5, 3.95, 3.23, 0.35, "局限（论文里能看到的）", 13, True, ORANGE)
    bullets(s, 9.5, 4.3, 3.23, p["limit"], 13, MUTE, 4)
    footer(s, "一句话带走：" + p["take"])

# ---------- 9 放在一起看 ----------
s = prs.slides.add_slide(BLANK)
header(s, "放在一起看", "五篇的“问题—做法—证据—局限”对照")
cols = ["论文", "要解决的问题", "核心做法", "最关键的证据", "已知局限"]
widths = [1.5, 2.6, 2.9, 2.6, 2.5]
rows = [[p["head"].split("：")[0], p["problem"][0], p["method"][0], p["stat"], p["limit"][0]] for p in N["papers"]]
tbl = s.shapes.add_table(len(rows)+1, len(cols), Inches(0.6), Inches(1.7), Inches(sum(widths)), Inches(4.6)).table
for j, w in enumerate(widths): tbl.columns[j].width = Inches(w)
def cell(c, t, size, bold=False, color=INK, fill=None):
    c.text = ""; p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = t; font(r, size, bold, color)
    c.vertical_anchor = MSO_ANCHOR.MIDDLE
    c.margin_left = c.margin_right = Inches(0.08)
    if fill is not None: c.fill.solid(); c.fill.fore_color.rgb = fill
for j, h in enumerate(cols): cell(tbl.cell(0, j), h, 13, True, WHITE, BLUE)
for i, r in enumerate(rows, 1):
    for j, t in enumerate(r): cell(tbl.cell(i, j), t, 12, j == 0, INK, SOFT2 if i % 2 == 0 else WHITE)
footer(s, "共同点：每篇都有一个清楚的反例或代价，不是只有成绩。")

# ---------- 10 对我们课题的启发 ----------
s = prs.slides.add_slide(BLANK)
header(s, "示例页：换成你自己的课题", "对我们课题的启发和下周计划")
items = [("启发 1", "基线：", "把预训练模型 + 线性探测作为新基线，先看我们的方法还剩多少空间。"),
         ("启发 2", "数据：", "ViT 和 CLIP 都说明规模和数据很关键；我们的数据量够不够，要先估一下。"),
         ("启发 3", "评估：", "几篇都在多个数据集上评估；我们除了本组数据，要加一个外部数据集。")]
for i, (a, b, c) in enumerate(items):
    y = 1.9 + i*1.15
    label(s, 0.6, y, 1.3, 0.75, a, 15, BLUE, WHITE, True)
    text(s, 2.1, y+0.05, 6.5, 0.75, [b + c], 16, False, INK, anchor=MSO_ANCHOR.MIDDLE)
box(s, 9.0, 1.9, 3.73, 3.35, fill=SOFT)
text(s, 9.2, 2.0, 3.4, 0.4, "下周计划", 16, True, BLUE)
bullets(s, 9.2, 2.5, 3.4, ["跑通预训练模型 + 线性探测基线", "整理一版数据说明", "找一个外部数据集", "下次组会汇报基线结果"], 14, INK, 8)
footer(s, "需要老师拍板的：基线用哪个预训练模型？外部数据集从哪里拿？")

# ---------- 11 备用页 ----------
s = prs.slides.add_slide(BLANK)
header(s, "备用页：老师可能会问", "提前备好的回答（来自论文原文）")
qa = [("ViT 为什么在 ImageNet 上不如 ResNet？", "没有强正则时，ViT 比同等大小的 ResNet 低几个百分点；论文认为缺少平移等变性、局部性等归纳偏置，数据不够时泛化差；预训练到 1400 万–3 亿张图后情况反转。"),
      ("ResNet 的 3.57% 是单模型吗？", "不是，是集成的 ILSVRC 2015 测试集 top-5 错误率；152 层单模型验证集 top-5 是 4.49%。"),
      ("CLIP 的零样本是怎么做的？", "把类别名写成文字，用文本编码器得到类别向量，和图像向量比相似度，选最像的。"),
      ("DDPM 的 FID 3.17 是在什么设置下？", "CIFAR-10，无条件生成；所有实验设 T = 1000。"),
      ("Transformer 的主要代价是什么？", "自注意力每层复杂度 O(n²·d)，序列长了会很贵（论文 Table 1）。")]
for i, (q, a) in enumerate(qa):
    y = 1.85 + i*0.92
    text(s, 0.6, y, 4.3, 0.85, q, 14, True, BLUE, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 5.0, y, 7.7, 0.85, a, 14, False, INK, anchor=MSO_ANCHOR.MIDDLE)
    ln = s.shapes.add_connector(1, Inches(0.6), Inches(y+0.88), Inches(12.73), Inches(y+0.88)); ln.line.color.rgb = LINE; ln.line.width = Pt(1); st = ln._element.find(qn('p:style')); ln._element.remove(st) if st is not None else None; ln.line.color.rgb = LINE

out = os.path.join(HERE, "组会示例.pptx"); prs.save(out); print("已生成", out, "共", len(prs.slides), "页")
